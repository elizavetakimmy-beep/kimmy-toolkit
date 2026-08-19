#!/usr/bin/env python3
"""
Robust loader + inspector for messy real-world spreadsheets (CSV / .xlsx).

Two ways to use it:

  1) Inspect the structure before parsing (ALWAYS do this first):
         python3 parse_spreadsheet.py --inspect "file.xlsx"
         python3 parse_spreadsheet.py --inspect "file.csv" --rows 40 --cols 20

  2) Import the helpers into your own small parser:
         from parse_spreadsheet import load, clean_number
         rows = load("file.csv")            # list[list[str|None]]  (CSV)
         sheets = load("file.xlsx")         # dict[name -> list[list]] (xlsx)

Why a script and not "just read the file": real sheets are wide (200+ cols),
have merged/repeated header rows, Plan/Fact/Diff column triples, month x week
layouts, and error cells. Reading the raw file into the model's context wastes
tokens and hides the structure. Inspect -> reason about layout -> parse with code.

The single most important helper is clean_number(): it turns the messy strings
people actually type ("1 234,5 ₽", "5%", "р.90 560", "#REF!", non-breaking
spaces) into floats or None. Missing/error -> None, NEVER 0 (0 is a real value
and pretending a blank is 0 corrupts every average and total downstream).
"""
import sys
import csv
import argparse

ERROR_MARKERS = ("#REF", "#DIV", "#VALUE", "#N/A", "#NAME", "#NULL", "#NUM")


def clean_number(v):
    """Normalize a spreadsheet cell to a float, or None if it isn't a number.

    Handles: currency (₽, р., руб, $), percent (%), thousands separators
    including the non-breaking/thin spaces Excel loves, comma decimals
    ("1 234,5" -> 1234.5), and formula errors (#REF! etc. -> None).
    Returns None for blanks and errors so downstream sums/means stay correct.
    """
    if v is None:
        return None
    s = str(v).strip()
    if s == "" or any(e in s for e in ERROR_MARKERS):
        return None
    # strip currency / percent / unit noise
    for junk in ("₽", "р.", "руб", "$", "€", "%"):
        s = s.replace(junk, "")
    # normalize spaces (regular, non-breaking, thin) and comma decimals
    s = s.replace("\xa0", " ").replace(" ", " ").replace(" ", "")
    s = s.replace(",", ".")
    if s in ("", "-", ".", "—", "–"):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def is_percent(v):
    """True if the raw cell was written as a percentage (had a % sign)."""
    return v is not None and "%" in str(v)


def load(path):
    """Load a CSV (-> list of rows) or .xlsx (-> dict of sheetname -> rows).

    Cells come back as raw strings/values; run them through clean_number()
    when you need numbers. xlsx uses openpyxl with data_only=True so formula
    cells return their last cached computed value.
    """
    low = path.lower()
    if low.endswith((".xlsx", ".xlsm", ".xltx")):
        import openpyxl  # pip install openpyxl if missing
        wb = openpyxl.load_workbook(path, data_only=True)
        out = {}
        for ws in wb.worksheets:
            out[ws.title] = [
                [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
                for r in range(1, ws.max_row + 1)
            ]
        return out
    # CSV / TSV
    delim = "\t" if low.endswith(".tsv") else ","
    with open(path, encoding="utf-8-sig") as f:
        return list(csv.reader(f, delimiter=delim))


def _fmt_cell(v):
    s = "" if v is None else str(v)
    s = s.replace("\n", "\\n")
    return s if len(s) <= 22 else s[:21] + "…"


def inspect(path, max_rows=30, max_cols=18):
    """Print a compact map of the file: sheets, dimensions, and a grid preview
    showing only non-empty cells with their (col, value). This is what you read
    to figure out the layout before writing a parser."""
    data = load(path)
    sheets = data if isinstance(data, dict) else {"(csv)": data}
    for name, rows in sheets.items():
        ncols = max((len(r) for r in rows), default=0)
        print(f"\n===== SHEET {name!r}  rows={len(rows)}  cols={ncols} =====")
        for ri, row in enumerate(rows[:max_rows]):
            cells = [
                f"c{ci}={_fmt_cell(v)!r}"
                for ci, v in enumerate(row[:max_cols])
                if v not in (None, "")
            ]
            if cells:
                print(f"  r{ri:<3} " + " | ".join(cells))
        if len(rows) > max_rows:
            print(f"  … ({len(rows) - max_rows} more rows)")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Inspect or load a spreadsheet.")
    ap.add_argument("file")
    ap.add_argument("--inspect", action="store_true",
                    help="print the structure (default action)")
    ap.add_argument("--rows", type=int, default=30, help="max rows to preview")
    ap.add_argument("--cols", type=int, default=18, help="max cols to preview")
    args = ap.parse_args()
    # --inspect is the only real action; keep the CLI dead simple.
    inspect(args.file, max_rows=args.rows, max_cols=args.cols)
